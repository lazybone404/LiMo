"""语义标记 `limo-block` 的**导出与下游消费点** —— 四个出口 + 两个数据层。

协议与解析本身的测试在 `test_artifact_blocks.py`（读共享 fixture，与前端逐条对称）。
**本文件测的是「谁消费它、消费得对不对」**——那是协议解析全绿也照样会出问题的一半。

本文件要防的失败**全都是静默的**（断言全过、肉眼等于没有的那一类）：

  A. HTML 导出漏掉 `<div class="md-block">` 包装 → 标记行**原样出现在导出页上**给读者
     看（这正是用户最初提的现象：「解析失败剥掉的标记…在前端显示」）。断「块内文字
     还在」是绿不了它的 —— 内容一直都在。
  B. Word 侧漏加 `div.md-block` 分支 → 落进 `_render_block` 末尾的 `else`「未知块：
     递归子元素」→ 内容一个不少，只是没有强调块形态。与 A 同理，必须断到**结构层**
     （表格数 / 底纹 / 左边框）。
  C. 无标记的产物也注入了 `blocks.css` 或写了 `data-marker-decor` → 「导出与改造前
     逐字节等价」这条承诺**静默失效**：没有一条断言会红，而每一份历史产物的导出 HTML
     都多了一段 CSS。
  D. md / 纯文本导出没剥标记 → `.md` / `.txt` 文件开头/中间留着一行
     `<!-- limo-block:v1 … -->`（纯文本里它是**可见噪音**）。
  E. 合成管道没剥标记 → 标记行进模型 prompt、且被 `estimate_tokens` 算进逐篇配额。
  F. `chars` 把标记行也算进去 → 「字数」这个数从此与读者看到的字数不一致。
"""
from __future__ import annotations

import io
import re
import shutil
import tempfile
import zipfile

from docx import Document
from docx.oxml.ns import qn
from fastapi import FastAPI
from fastapi.responses import JSONResponse
from fastapi.testclient import TestClient

from app.core.artifact_blocks import ARTIFACT_BLOCK_TYPES, strip_blocks
from app.core.dependencies import get_store
from app.core.errors import AppError
from app.core.transcript_meta import strip_meta
from app.core.types import (GenerateMode, MarkdownView, RawMetadata, RawTranscript,
                            SourceType, Style)
from app.generation.doc_text import latest_notes_text
from app.routers import export as E
from app.routers.export import _build_selfcontained_html, _load_block_css
from app.services.docx_export import BAR_MARKER, FILL_MARKER, markdown_to_docx
from app.storage.store import Store

_OPEN = '<!-- limo-block:v1 type="key-takeaway" -->'
_CLOSE = '<!-- /limo-block -->'
BLOCK = f"{_OPEN}\n**核心结论**\n\n采用异步任务处理长文档，可以避免请求超时。\n{_CLOSE}"

MD = f"开场白。\n\n## 甲\n内容一\n\n{BLOCK}\n\n## 乙\n内容二\n"
PLAIN = "开场白。\n\n## 甲\n内容一\n\n## 乙\n内容二\n"

#: 卡片轴的用例专用（2026-10-11 起**只有被 `card` 节修饰符标过的节**才成卡，
#: 故 `MD` 那份点不出一张卡来）。⚠️ `MD` 保持无标记不动 —— 它被本文件里其余
#: 十几个用例消费，给它加标记等于顺手改掉一批与骨架无关的用例的输入。
CARD_LINE = '<!-- limo-block:v1 type="card" -->'
CARD_MD = (
    f"开场白。\n\n## 甲\n{CARD_LINE}\n内容一\n\n{BLOCK}\n\n"
    f"## 乙\n{CARD_LINE}\n内容二\n"
)


# ═══════════════════════════════════════════════════════════════
# 0. 夹具：一份**含标记**的 notes 产物（走真实 Store，非桩）
# ═══════════════════════════════════════════════════════════════

_TMP: list[str] = []


def _new_store() -> Store:
    d = tempfile.mkdtemp()
    _TMP.append(d)
    return Store(d)


_STORE = _new_store()
_STORE.save_transcript(RawTranscript(
    id='tr_blk001', title='标记样例', source_type=SourceType.WEBPAGE,
    engine='trafilatura', raw_text='# 原文\n\n正文', content_hash='blk1',
    metadata=RawMetadata(),
))
_STORE.save_markdown(MarkdownView(
    id='md_blk_notes', transcript_id='tr_blk001', content=MD,
    style=Style.DETAILED, mode=GenerateMode.NOTES,
))

#: 证据块夹具（2026-10-18，多篇合成专用）。两篇：
#: `tr_ev001` 只有证据块（用来证明「线性化后已无块 ⇒ 不注入 blocks.css」），
#: `tr_evmix1` 证据块 + 强调块共存（用来证明只消解 evidence、其余块照常渲染）。
_EV_OPEN = '<!-- limo-block:v1 type="evidence" -->'
_EV_CLOSE = '<!-- /limo-block -->'
EVIDENCE_MD = (
    "论断：A 优于 B [文档 2]。\n\n"
    f"{_EV_OPEN}\n- 样本：N=120，跨 3 年\n- 局限：单一中心\n{_EV_CLOSE}\n\n"
    "后续正文。\n"
)
EV_MIXED_MD = EVIDENCE_MD + "\n" + BLOCK + "\n"

for _tid, _md, _hash in (
    ('tr_ev001', EVIDENCE_MD, 'ev1'),
    ('tr_evmix1', EV_MIXED_MD, 'ev2'),
):
    _STORE.save_transcript(RawTranscript(
        id=_tid, title='证据样例', source_type=SourceType.WEBPAGE,
        engine='trafilatura', raw_text='# 原文\n\n正文', content_hash=_hash,
        metadata=RawMetadata(),
    ))
    _STORE.save_markdown(MarkdownView(
        id=f'md_{_tid}', transcript_id=_tid, content=_md,
        style=Style.DETAILED, mode=GenerateMode.NOTES,
    ))

app = FastAPI()
app.include_router(E.router, prefix="/api")
app.dependency_overrides[get_store] = lambda: _STORE


@app.exception_handler(AppError)
async def _app_error_handler(request, exc: AppError):
    # ⚠️ 与 main.py **同规则**（模式 #36）：显式 status_code 优先，未显式才 recoverable→200。
    # 与 `test_export_custom_zip.py` 逐字同源 —— 自造一份更宽松的契约会让判据失去意义。
    http_status = exc.status_code if exc._explicit_status else (
        200 if exc.recoverable else exc.status_code)
    return JSONResponse(status_code=http_status, content=exc.to_dict())


client = TestClient(app)


def _export(**kw) -> dict[str, bytes]:
    # `doc_ids` 允许覆盖（证据块那组用例消费的是另一篇夹具文档），默认保持原样 ——
    # 本文件前面十几条用例的输入不该因为新增一组用例而改变。
    doc_ids = kw.pop('doc_ids', ['tr_blk001'])
    body = {'doc_ids': doc_ids, 'content_tabs': ['summary'], **kw}
    r = client.post('/api/export/custom', json=body)
    assert r.status_code == 200, f"{r.status_code} {r.text[:300]}"
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        return {n: zf.read(n) for n in zf.namelist() if not n.endswith('/')}


def _only(files: dict[str, bytes]) -> str:
    assert len(files) == 1, files.keys()
    return next(iter(files.values())).decode('utf-8')


def _html_tag(html: str) -> str:
    """`<html …>` 开标签本身 —— 装饰开关是写在这上面的属性。

    ⚠️ **不能拿整篇断子串**：内联的 `blocks.css` 里那些注释逐字写着
    `data-marker-decor="off"` / `limo-block`，于是 `'…' in html` 恒真、
    `'…' not in html` 恒假，两个方向都测不到东西。
    """
    return html[:html.index('>', html.index('<html')) + 1]


def _body(html: str) -> str:
    """`<body>` 之后的正文 —— 断言「标记行没露到读者眼前」必须排除 `<style>` 块：
    内联的 `blocks.css` **注释**里逐字写着示例标记（`<!-- limo-block:v1 type="…" -->`），
    拿整篇断 `'<!-- limo-block' not in html` 会恒假。
    """
    return html[html.index('<body>'):]


# ═══════════════════════════════════════════════════════════════
# A. HTML 链
# ═══════════════════════════════════════════════════════════════

def test_无标记_html_导出逐字节不变():
    """零回归：**没有标记的产物**导出 HTML 与「开关不存在」的改造前完全一样。

    对照 C：把注入闸门写成「无条件」时本条红 —— 而后果是每一份历史产物的导出 HTML
    都凭空多一段 CSS，没有任何断言会红、也没有任何用户会察觉（样式恰好没生效）。
    """
    base = _build_selfcontained_html(PLAIN, "标题")
    assert base == _build_selfcontained_html(PLAIN, "标题", marker_decor=False)
    assert base == _build_selfcontained_html(PLAIN, "标题", marker_decor=True)
    # ⚠️ 判据**必须精确到 class 属性**，不能用裸子串 `'md-block' in html`：
    # limo 主题自 2026-10-16 起有 `--md-blockquote-border` / `--md-blockquote-bg`，
    # 它**含**这个子串 —— 裸子串版本当场变红，而它红的原因是**判据太宽**，
    # 与「真有块漏包」完全无关。`_body` 再排除 `<style>` 与注释里的示例。
    assert 'class="md-block' not in _body(base)
    assert 'data-marker-decor' not in base


def test_有标记_html_块被包成_div_而不是裸露的标记行():
    """对照 A：漏掉包装时标记行会**原样出现在页面上**给读者看。

    ⚠️ 同时断两件事：包装**在**（结构），标记行**不在**（症状）。只断前者的话
    「包了一层但仍把注释拼进正文」照样绿。
    """
    html = _build_selfcontained_html(MD, "标题")
    assert html.count('<div class="md-block" data-block-type="key-takeaway">') == 1
    # ⚠️ 判据是「**读者可见的正文里**没有标记行」而不是整篇子串 —— 内联的 blocks.css
    # 注释里必然含这个协议的名字与示例（那份文件通篇在讨论它）。
    assert '<!-- limo-block' not in _body(html)
    assert '<!-- /limo-block' not in _body(html)
    # 块内文字一字不少，块外正文也没丢
    for chunk in ("核心结论", "采用异步任务处理长文档", "开场白。", "内容一", "内容二"):
        assert chunk in html, chunk


def test_有标记_html_内联了_blocks_css():
    """⚠️ 断的是「导出页里出现了 `blocks.css` 的选择器」，不是「文件存在」——
    文件存在但没拼进 style 块，效果与没有一样（与卡片那条同款）。"""
    html = _build_selfcontained_html(MD, "标题")
    css = _load_block_css()
    assert '[data-marker-decor="off"]' in html
    assert '[data-marker-decor="off"]' in css      # 自证：这份 CSS 确实是标记规则那份
    assert '.md-block' in html


def test_blocks_css_引用的每个界面令牌在导出页里都有定义():
    """**成对常量的守卫**：`blocks.css` 里每处 `var(--tk-…)` 都必须在导出 HTML 里有定义。

    ⚠️ 导出主题（`HTML_THEMES`）是自包含的独立 CSS，**一个 `--tk-*` 都不提供** ——
    `var()` 无 fallback 时整条声明作废。漏一个的表现是「强调块少一块样式」，不报错、
    不抛异常，且**应用内预览里是好的**（那边令牌由主题系统提供）。这条是唯一能自动
    发现「blocks.css 新引了一个令牌而导出侧没补」的地方。
    """
    import re as _re
    css = _re.sub(r"/\*[\s\S]*?\*/", "", _load_block_css())   # 剥注释：注释里提到名字不算引用
    used = sorted(set(_re.findall(r"var\((--tk-[a-z-]+)", css)))
    assert used, "blocks.css 没引用任何 --tk-* —— 判据本身失效了（改名了？）"
    html = _build_selfcontained_html(MD, "标题")
    missing = [t for t in used if f"{t}:" not in html]
    assert missing == [], f"导出页缺少这些令牌的定义：{missing}"


def test_装饰开关_关时写_off_开时不写属性():
    """«关» 必须是**属性存在且为 off**（CSS 靠它归零装饰）；«开» 时属性**整个不存在**，
    而不是等于 `"on"` —— 后者把「零回归」从「选择器不匹配」退化成「取值相同」。
    """
    off = _build_selfcontained_html(MD, "标题", marker_decor=False)
    on = _build_selfcontained_html(MD, "标题", marker_decor=True)
    assert 'data-marker-decor="off"' in _html_tag(off)
    assert 'data-marker-decor' not in _html_tag(on)
    # ⚠️ 关掉装饰**不改变内容**（DOM 保留、只是 CSS 归零）—— 一字不少
    for chunk in ("核心结论", "采用异步任务处理长文档"):
        assert chunk in off, chunk


def test_card_骨架下块落在卡内():
    """骨架轴 ⊥ 标记轴：`card` 下强调块在**卡片内部**，而不是另起一节或被切卡逻辑吞掉。"""
    html = _build_selfcontained_html(CARD_MD, "标题", skeleton="card")
    assert html.count('<section class="md-card">') == 2      # 甲 / 乙（都标了 card）
    i_card = html.index('<section class="md-card">')
    i_block = html.index('<div class="md-block"')
    assert i_block > i_card, "强调块跑到卡片外面去了"
    assert html.count('<div class="md-block"') == 1


def test_块内的_H2_不会被切成新卡():
    """标记块里的 `##` 不是文档结构 —— 切卡必须跳过块内（与前端 `markdown-scan` 同源）。"""
    md = (f"开场\n\n## 甲\n{CARD_LINE}\n内容\n\n"
          '<!-- limo-block:v1 type="key-takeaway" -->\n## 这不是标题\n块内\n<!-- /limo-block -->\n')
    html = _build_selfcontained_html(md, "标题", skeleton="card")
    assert html.count('<section class="md-card">') == 1, "块内的 ## 被当成了新章节"


def _md_of(type_: str, body: str = "块内内容") -> str:
    """单个指定 type 的标记块 —— 新增类型逐一过导出链时用（`MD` 那份只含 key-takeaway）。"""
    return f'<!-- limo-block:v1 type="{type_}" -->\n{body}\n<!-- /limo-block -->\n'


#: `steps` / `timeline` 的**形状**全靠块内是 `<ul>` 还是 `<ol>`（`blocks.css` 按子元素
#: 类型区分，见文件头），故这两个的用例必须带列表 body —— 用纯文字 body 的话，
#: 「块内列表结构原样穿过导出链」这件事根本没有证据，而它正是纯 CSS 方案的前提。
_LIST_BODY = {
    "steps": "- 输入\n- 清洗\n- 输出",
    "timeline": "1. **导入**\n2. **解析**",
}


def test_html_新增八种类型都包成_md_block_而不是裸露的标记行():
    """**「解析器对 type 无感知」这条设计在导出侧的证据**：八种新形态在 Python 侧
    一行分支都没有，走的正是 key-takeaway 的同一条包装路径，形态差异全在 CSS。

    ⚠️ 断两件事：包装在（结构）+ 标记行不在读者可见正文里（症状）。只断前者的话
    「包了一层但仍把注释拼进正文」照样绿 —— 与 `key-takeaway` 那条同款。
    """
    for t in ("alert", "recommend", "metric", "steps", "timeline",
              "table", "code", "image"):
        html = _build_selfcontained_html(_md_of(t), "标题")
        assert f'<div class="md-block" data-block-type="{t}">' in html, t
        assert '<!-- limo-block' not in _body(html), t
        assert '<!-- /limo-block' not in _body(html), t
        assert '块内内容' in html, t


#: `table` / `code` / `image` 的**全部增量就是一层外壳** —— 块内那张表、那段代码、
#: 那张图必须**原样**穿过导出链。用默认的纯文字 body 测不出这一点（文字放哪儿都是
#: 文字），故这三个必须带专属 body。与 `_LIST_BODY` 同一条理由：纯 CSS 方案的前提是
#: 「Markdown 结构在链路上不被压扁」，而那件事没有断言就只是一句愿望。
_RICH_BODY = {
    "table": "| 方案 | 成本 |\n| --- | --- |\n| A | 低 |",
    "code": "```python\ncfg = load()\n```",
    "image": "![架构图](assets/arch.png)",
}


def _text_of(html: str) -> str:
    """HTML → 纯文字（只为断言「内容还在」，不解码实体 —— 本文件这几条用例里没有实体）。

    ⚠️ 非有它不可的原因：Pygments 会把 `cfg = load()` 拆成一串 `<span>`，直接
    `'cfg = load()' in html` 会**假红**（代码明明在，只是被逐 token 高亮了）。断言
    「文字在不在」本就不该受高亮实现的影响。
    """
    return re.sub(r"<[^>]+>", "", html)


def test_html_内容卡片三种_块内富结构原样穿过导出链():
    """这一族**唯一会坏的方式**不是「块没包上」，而是**块内那份富内容被压扁** ——
    表格变成一行字、代码变成一行字、图片变成一个失效引用。那时 `data-block-type`
    包装断言照样绿、文字也一字不少（因为「压扁」压掉的是结构不是文字）。

    ⚠️ 这三个的整个价值就是「内容还是原来那份内容」，故这条断言与上面那条包装断言
    缺一不可 —— 上面管壳，这条管内容。
    ⚠️ **两层分开断**：标签归标签（结构），文字归文字（内容）。混在一起断就会出现
    「代码在、只是被高亮了」这种假红。
    """
    body = _body(_build_selfcontained_html(
        "\n".join(_md_of(t, b) for t, b in _RICH_BODY.items()), "标题"))
    # 结构层：三样富内容各自的标签都真的在
    assert "<table>" in body, "块内表格没有原样穿过（或没被渲染成 <table>）"
    assert "<pre" in body and "<code" in body, "块内代码没有原样穿过"
    # ⚠️ 图片的说明文字在 **`alt` 属性**里（不在文本节点），故它归结构层 ——
    # 剥标签时属性会一起被剥掉，写成 `'架构图' in _text_of(body)` 会假红。
    assert "<img" in body and "arch.png" in body, "块内图片没有原样穿过"
    assert 'alt="架构图"' in body, "图片的说明文字（alt）丢了"
    # 内容层：文字一个不少（剥掉标签后再断，免得被高亮的 span 拆解骗过）
    text = _text_of(body)
    assert "方案" in text and "低" in text, "表格里的文字丢了"
    assert "cfg = load()" in text, "代码里的文字丢了"


def test_html_列表结构原样穿过导出链():
    """`steps` / `timeline` 的形态判据是**块内子元素类型**（无序 / 有序列表），而
    Python 侧对此一行分支都没有 —— 形状完全靠 Markdown 渲染器把 `- ` / `1. ` 转成
    `<ul>` / `<ol>`。若哪一步（KaTeX 保护链 / 骨架切分 / 段拼接）把列表压成纯文本，
    `blocks.css` 里那两条按子元素选择器写的规则就**静默不命中**，两个 type 一起退化成
    裸段落 —— 而包装断言 `data-block-type="steps"` 照样绿。

    ⚠️ 这是「CSS 单一真相源」这套方案**唯一依赖的 Markdown 层前提**，值得单独钉住。
    """
    h_steps = _body(_build_selfcontained_html(_md_of("steps", _LIST_BODY["steps"]), "标题"))
    h_tl = _body(_build_selfcontained_html(_md_of("timeline", _LIST_BODY["timeline"]), "标题"))
    # 对照：steps 用无序列表、timeline 用有序列表 —— 两边**各自**只出自己要的那一种
    # （只断「有 <ul>」的话，一个把两者都渲染成 <ul> 的实现照样绿）。
    assert "<ul>" in h_steps and "<ol>" not in h_steps, "steps 应渲染成无序列表"
    assert "<ol>" in h_tl and "<ul>" not in h_tl, "timeline 应渲染成有序列表"


# ═══════════════════════════════════════════════════════════════
# B. Word 链
# ═══════════════════════════════════════════════════════════════

def _tables(data: bytes) -> list:
    return Document(io.BytesIO(data)).tables


def _paragraph_texts(data: bytes) -> list[str]:
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs)
    return texts


def test_word_无标记_零表格():
    """零回归：没有标记时 Word 结构与改造前一致（改造前也不产表格）。"""
    assert len(_tables(markdown_to_docx(PLAIN, "标题"))) == 0


def test_word_有标记_块包成单单元格表格():
    """对照 B：断在**结构层**（表格数 / 单行单列），不断「内容还在」——
    漏加 `div.md-block` 分支时内容由 `else` 兜底照旧渲染，那种断言绿得毫无意义。"""
    tables = _tables(markdown_to_docx(MD, "标题"))
    assert len(tables) == 1
    assert len(tables[0].rows) == 1
    assert len(tables[0].columns) == 1


def test_word_标记块的底纹与左边框真的生效():
    """结构层再深一层：只有表格没有底纹/竖线 = 一个白框（读者看不出这是重点）。"""
    tables = _tables(markdown_to_docx(MD, "标题"))
    tc = tables[0].rows[0].cells[0]._tc
    tcPr = tc.find(qn("w:tcPr"))
    assert tcPr is not None
    shd = tcPr.find(qn("w:shd"))
    assert shd is not None and shd.get(qn("w:fill")) == FILL_MARKER
    borders = tcPr.find(qn("w:tcBorders"))
    assert borders is not None, "缺少单元格左边框（强调块的关键视觉）"
    left = borders.find(qn("w:left"))
    assert left is not None and left.get(qn("w:color")) == BAR_MARKER


def test_word_标记行不出现在文档正文里():
    """**修的是现存泄漏缺陷**：BeautifulSoup 把 `<!-- … -->` 解析成 `Comment`
    （`NavigableString` 子类、**不是 Tag**）→ 走「非 Tag 即文本」的兜底分支 →
    标记被当成可见文字排成一个段落。这条断言就是那个缺陷的回归守卫。
    """
    texts = "\n".join(_paragraph_texts(markdown_to_docx(MD, "标题")))
    assert 'limo-block' not in texts
    assert '<!--' not in texts


def test_word_块内文字一字不少():
    """结构对了还要内容对 —— 不能把块内文字吞掉或复制一份。"""
    texts = "\n".join(_paragraph_texts(markdown_to_docx(MD, "标题")))
    assert texts.count("核心结论") == 1
    assert texts.count("采用异步任务处理长文档，可以避免请求超时。") == 1
    for chunk in ("开场白。", "内容一", "内容二"):
        assert chunk in texts, chunk


def test_word_card_骨架下标记块表格嵌在卡片表格内():
    """两个轴同时开：顶层 2 张卡，标记块那只表格**嵌在卡片里面**。

    ⚠️ `Document.tables` 只返回 **body 顶层**的表格 —— 嵌套表在 `cell.tables` 里。
    写成「顶层应该有 3 张」会红在这里，而实现其实是对的。
    """
    doc = Document(io.BytesIO(markdown_to_docx(CARD_MD, "标题", skeleton="card")))
    assert len(doc.tables) == 2, "卡片数不对（甲 / 乙）"
    nested = [inner for t in doc.tables
              for row in t.rows for cell in row.cells for inner in cell.tables]
    assert len(nested) == 1, "标记块没有嵌在卡片内部（或反过来，卡片嵌在块里）"
    assert 'limo-block' not in "\n".join(
        p.text for p in nested[0].rows[0].cells[0].paragraphs)


def _marker_colors(data: bytes) -> tuple[str, str]:
    """从 Word 产物里**读回**强调块的 (底纹色, 左边框色)。

    ⚠️ 刻意不 `import _MARKER_COLORS` 来比对 —— 那样只能证明「表里写着什么」，
    证明不了「渲染时真查了表」。漏查表时的表现正是**全部走兜底色**：表本身完全正确，
    产物上三种块长得一模一样，而 import 式的断言一路绿。
    """
    tables = _tables(data)
    assert len(tables) == 1, f"强调块没有包成单张表格（{len(tables)} 张）"
    tcPr = tables[0].rows[0].cells[0]._tc.find(qn("w:tcPr"))
    shd = tcPr.find(qn("w:shd"))
    left = tcPr.find(qn("w:tcBorders")).find(qn("w:left"))
    return shd.get(qn("w:fill")), left.get(qn("w:color"))


def test_word_各类型的配色互不相同():
    """加了块却只有一种视觉 = 白加。这条挡的是「`_render_marker_block` 漏按
    `data-block-type` 取色」——那种情况下六种 type 全走 `(FILL_MARKER, BAR_MARKER)`。

    ⚠️ 断**两两不同**而不是逐个比对写死的色值：写死色值等于把实现抄进测试，
    改配色要改两处、改错一处时测试还是绿的。
    ⚠️ 名单从 `_MARKER_COLORS` 的键派生（`import` 的是**键集**不是**值**）——
    手写名单在下次加 type 时会绿着失去意义（模式 #54）。上面 `_marker_colors` 那条
    「不 import 值比对」的纪律不受影响：这里仍然是从产物**读回**颜色再两两比。
    """
    from app.services.docx_export import _MARKER_COLORS
    types = sorted(_MARKER_COLORS)
    got = {t: _marker_colors(markdown_to_docx(_md_of(t), "标题")) for t in types}
    assert len(set(got.values())) == len(types), \
        f"有类型共用同一套配色（漏查表的症状）：{got}"


def test_word_新增类型的块内文字一字不少():
    """结构对了还要内容对 —— 不能因为加分支把块内文字吞掉或复制一份。

    ⚠️ 名单从 `ARTIFACT_BLOCK_TYPES` **派生**（排除 `key-takeaway` —— 它由上面 `MD`
    那条覆盖），手写名单在下次加 type 时会绿着失去意义（模式 #54）。
    """
    for t in sorted(ARTIFACT_BLOCK_TYPES - {"key-takeaway"}):
        texts = "\n".join(_paragraph_texts(markdown_to_docx(_md_of(t), "标题")))
        assert texts.count("块内内容") == 1, t
        assert 'limo-block' not in texts, t


def test_word_内容卡片三种_块内富结构原样穿过且_table_是表格套表格():
    """`table` / `code` / `image` 的价值同样是「块内那份富内容原样保留」，Word 侧要断
    两件事：

      · **内容在** —— 表格的行列、代码的文字、图片的占位都在（不是只有一个空壳）；
      · **形态边界** —— `table` 卡在这里是**表格套表格**：外层是强调块用的单单元格
        表格，内层是 Markdown 表格自己那张 Table Grid。⚠️ 这是**能力边界不是缺陷**：
        Word 的「容器」手段只有表格这一种，要不嵌套就得让 Python 去识别「块内是什么」，
        而那正是 `blocks.css` 文件头否掉的方向（形态差异只走 CSS）。钉住现状，不修。

    ⚠️ `Document.tables` 只返回 **body 顶层**的表格，嵌套表在 `cell.tables` 里 ——
    写成「顶层应该有 2 张」会红在这里，而实现其实是对的（与
    `test_word_card_骨架下标记块表格嵌在卡片表格内` 同款陷阱）。
    """
    doc = Document(io.BytesIO(markdown_to_docx(_md_of("table", _RICH_BODY["table"]), "标题")))
    assert len(doc.tables) == 1, "table 卡的外壳没有包成单单元格表格"
    shell = doc.tables[0]
    assert len(shell.rows) == 1 and len(shell.columns) == 1, "外壳应是一格"
    nested = shell.rows[0].cells[0].tables
    assert len(nested) == 1, "块内那张表格没有原样穿过（或被压扁成文字）"
    inner = "\n".join(p.text for r in nested[0].rows for c in r.cells for p in c.paragraphs)
    assert "方案" in inner and "低" in inner, "内层表格的内容丢了"

    # code / image：同样包成单格，块内的文字 / 图片占位一字不少
    for t, want in (("code", "cfg = load()"), ("image", "架构图")):
        texts = "\n".join(_paragraph_texts(markdown_to_docx(_md_of(t, _RICH_BODY[t]), "标题")))
        assert want in texts, f"{t} 丢了块内内容：{want}"
        assert 'limo-block' not in texts, t


def test_word_steps_timeline_退化成竖排列表_这是能力边界():
    """⚠️ **如实记录一条已知偏差**：`steps`（横向步骤条）在 Word 里**不是**横排 ——
    预览侧的横排靠 `display:flex` + 伪元素箭头，Word 的 XML 没有对应表达。两者在这里
    都渲染成一个单单元格表格 + 竖排列表，**只有配色不同**。

    把它写成断言而不是只写进注释：哪天有人以为「Word 侧也是横的」而按此改代码，
    这条会先红。修法（改多单元格表格）等于让 Python 复刻预览的结构 —— 那正是
    `blocks.css` 文件头明确否掉的方向，故**不修**，只钉住现状。
    """
    want = {"steps": ("输入", "清洗"), "timeline": ("导入", "解析")}
    for t in ("steps", "timeline"):
        doc = Document(io.BytesIO(markdown_to_docx(_md_of(t, _LIST_BODY[t]), "标题")))
        assert len(doc.tables) == 1, f"{t} 应包成单单元格表格"
        assert len(doc.tables[0].rows) == 1 and len(doc.tables[0].columns) == 1, t
        texts = "\n".join(_paragraph_texts(markdown_to_docx(_md_of(t, _LIST_BODY[t]), "标题")))
        # 列表项文字一字不少（形态退化、内容不退化）
        for w in want[t]:
            assert w in texts, f"{t} 丢了列表项文字：{w}"
        assert 'limo-block' not in texts, t


# ═══════════════════════════════════════════════════════════════
# C. 定制导出端点（md / plain / html / docx 四个格式各走一遍真链路）
# ═══════════════════════════════════════════════════════════════

def test_md_导出剥掉标记行且保留内容():
    """对照 D：`.md` 文件里留着 `<!-- limo-block:v1 … -->` —— 用户拍板「导出的文件应
    当是干净正文」。⚠️ 剥的是**标记行**，块内文字一个字不能少。"""
    out = _only(_export(formats=['md']))
    assert 'limo-block' not in out
    assert '核心结论' in out
    assert '采用异步任务处理长文档' in out
    assert '开场白。' in out


def test_plain_导出剥掉标记行且保留内容():
    """纯文本里 HTML 注释是**可见噪音**（读者会看到一整行 `<!-- … -->`），与 md 同理。"""
    out = _only(_export(formats=['plain']))
    assert 'limo-block' not in out
    assert '<!' not in out
    assert '核心结论' in out


def test_html_端点导出带强调块():
    out = _only(_export(formats=['html']))
    assert '<div class="md-block" data-block-type="key-takeaway">' in out
    assert '<!-- limo-block' not in _body(out)
    # ⚠️ paper（默认骨架）下 `--tk-*` 令牌也必须注入 —— 这条是「只有 skeleton=card 才注入」
    # 那个缺陷的回归守卫（表现：强调块在导出页上一点样式都没有，预览里却正常）。
    assert '--tk-accent:' in out


def test_html_端点跟随装饰开关():
    """预览与导出共用同一个开关（「用户在预览栏看到什么，导出就应该得到什么」）。"""
    off = _only(_export(formats=['html'], marker_decor=False))
    on = _only(_export(formats=['html'], marker_decor=True))
    assert 'data-marker-decor="off"' in _html_tag(off)
    assert 'data-marker-decor' not in _html_tag(on)


def test_docx_端点导出块是表格():
    """Word 走的是**转换**（不是剥离）：块在 .docx 里应当是「有边框底纹的块」。

    ⚠️ 走的是端点的默认参数（`marker_decor` 对它无意义 —— Word 没有装饰开关），
    与上面直接调 `markdown_to_docx` 的那组**成对**：一组测转换本身，一组测它真的
    被端点用上了（`_build_export_zip` 的 docx 分支传的是原始 `content` 而非 `processed`）。
    """
    blob = next(iter(_export(formats=['docx']).values()))
    doc = Document(io.BytesIO(blob))
    assert len(doc.tables) == 1
    assert 'limo-block' not in "\n".join(p.text for p in doc.paragraphs)


# ═══════════════════════════════════════════════════════════════
# D. 合成管道 —— 标记不进模型 prompt、不计 token 配额
# ═══════════════════════════════════════════════════════════════

def test_合成读原料时剥掉标记行():
    """对照 E：标记行进 prompt 会同时踩三条 —— 噪音、被 `estimate_tokens` 算进逐篇配额、
    材料里多一类我们专属的可控文本。⚠️ 块内文字**照旧进**（剥的只是标记行）。"""
    text = latest_notes_text(_STORE, 'tr_blk001')
    assert 'limo-block' not in text
    assert 'limo-meta' not in text        # 自证头成对地也剥（现有行为，别被这次改动带坏）
    assert '核心结论' in text              # 块内文字照旧
    assert '采用异步任务处理长文档' in text


# ═══════════════════════════════════════════════════════════════
# E. 字数统计 —— 标记行不算「正文字数」
# ═══════════════════════════════════════════════════════════════

def test_chars_不把标记行算进去():
    """对照 F：`chars` 是「正文字数」，标记行是读者看不到的协议元素。

    ⚠️ 判据写成「差值**恰好**等于两行标记 + 各自的行尾换行」而不是「chars < 原文长度」——
    后者在「只剥了开标记、漏了闭标记」时照样绿。
    """
    entry = _STORE.get_markdowns_by_transcript('tr_blk001', mode='notes')[0]
    raw = strip_meta(entry['content'])
    assert 'limo-block' in raw, "夹具失效：产物里没有标记，这条测不到东西"
    assert entry['chars'] == len(strip_blocks(raw))
    assert len(raw) - entry['chars'] == len(_OPEN) + 1 + len(_CLOSE) + 1


# ═══════════════════════════════════════════════════════════════
# F. 证据块线性化（合成专用类型，**导出四入口统一降级**）
# ═══════════════════════════════════════════════════════════════
#
# 判据本身（编号、来源、幂等、围栏内零触碰）在 `test_artifact_blocks.py`，读的是与
# 前端共享的 `evidence-linearize.json`。**本节测的是「四个入口谁真的调了它」** ——
# 协议层全绿而某一个入口漏接线，是本轮最容易漏、且**完全静默**的一类失败：
# 那种情况下导出成功、文件正常，只是折叠的内容在 md 里彻底消失。

def test_证据块_md_导出_原位占位加文末附录():
    """用户拍板「所有格式统一，Markdown 也一样」：md 里绝不能只剩折叠内容被吞掉。"""
    out = _only(_export(formats=['md'], doc_ids=['tr_ev001']))
    assert '[证据 1]' in out
    assert '## 证据附录' in out
    # 来源取自块**之前**那段正文的锚点（模型把来源写在论断旁，块内只写方法/样本）
    assert '**证据 1**（文档 2）' in out
    assert '样本：N=120' in out, "证据内容没有搬进附录"
    assert 'limo-block' not in out, "标记行漏进了 .md 文件"
    # 原位占位必须在附录之前 —— 顺序反了读者顺着 `[证据 1]` 往下翻什么也找不到
    assert out.index('[证据 1]') < out.index('## 证据附录')


def test_证据块_plain_导出同样线性化():
    """纯文本里折叠内容同样**没有**任何承载方式 —— 不线性化等于直接丢内容。"""
    out = _only(_export(formats=['plain'], doc_ids=['tr_ev001']))
    assert '[证据 1]' in out
    assert '## 证据附录' in out
    assert '样本：N=120' in out
    assert '<!' not in out


def test_证据块_html_导出_是纯文本引用而非折叠元素():
    """HTML 侧**刻意不还原成 `<details>`**：导出文件没有主题系统与交互，折叠 =
    默认收起 = 读者看不到依据（且打印/PDF 也拿不到）。线性化后 HTML 与其它格式读到的
    是同一份内容。"""
    body = _body(_only(_export(formats=['html'], doc_ids=['tr_ev001'])))
    assert '[证据 1]' in body
    assert '证据附录' in body
    assert '样本：N=120' in body
    assert '<details' not in body, "导出侧不该有折叠元素（预览侧才折叠）"


def test_只有证据块时导出不注入_blocks_css():
    """线性化后文档里**已经没有块**了 → `_has_blocks` 为假 → 不注入 `blocks.css`。

    ⚠️ 这条不只是「省一段 CSS」：`_has_blocks` 若为保持为真而偷偷留着围栏，HTML 导出
    会把标记行原样吐给读者（正是本文件开头 A 类失败）。判据与内容必须同时干净。
    """
    html = _only(_export(formats=['html'], doc_ids=['tr_ev001']))
    assert 'data-marker-decor' not in html
    assert 'class="md-block' not in _body(html)


def test_证据块与强调块共存时只有证据被消解():
    """同一个文档里两种块并存：证据被线性化，强调块照常包成 `md-block`。

    ⚠️ 这条挡的是「一刀切」实现：把整篇的 `has_blocks` 当开关，要么全消解要么全保留。
    两种情况都有一部分读者看到协议标记或丢内容，且各自只在一半文档上发作。
    """
    html = _only(_export(formats=['html'], doc_ids=['tr_evmix1']))
    assert '<div class="md-block" data-block-type="key-takeaway">' in html
    assert '[证据 1]' in html
    assert 'class="md-block" data-block-type="evidence"' not in html
    # 还有块 ⇒ blocks.css 照常注入（上面那条是「一个块都不剩」才不注入）
    assert 'data-marker-decor' in html


def test_证据块_docx_导出也拿到附录():
    """⚠️ Word 分支用的是**原始 content**（不是 `processed`，图片内嵌要原始引用）——
    线性化必须改 `content` 本身。只改 `processed` 时 md / html / plain 三条全对、
    Word 静默错（拿到含围栏的原文，或者对已经线性化的正文再剥一遍标记）。"""
    blob = next(iter(_export(formats=['docx'], doc_ids=['tr_ev001']).values()))
    texts = "\n".join(_paragraph_texts(blob))
    assert '[证据 1]' in texts
    assert '证据附录' in texts
    assert '样本：N=120' in texts
    assert 'limo-block' not in texts


def test_单文档导出端点也线性化():
    """`GET /export/{id}`（ExportMenu 的快捷导出 + PDF 两条路径的共同上游）。"""
    r = client.get('/api/export/tr_ev001', params={'format': 'md'})
    assert r.status_code == 200, r.text[:300]
    assert '[证据 1]' in r.text
    assert '## 证据附录' in r.text
    assert 'limo-block' not in r.text


def test_导出全部产物端点也线性化():
    """`GET /export/{id}/all` 走 `_collect_doc_products` —— 与 `/export/batch` **同一个**
    收集函数，故这一条同时覆盖两条路由。"""
    r = client.get('/api/export/tr_ev001/all')
    assert r.status_code == 200, r.text[:300]
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        names = [n for n in zf.namelist() if n.endswith('-笔记.md')]
        assert names, f"压缩包里没有笔记产物：{zf.namelist()}"
        text = zf.read(names[0]).decode('utf-8')
        # ⚠️ 转录原文**不该**被线性化（它不是产物，本来就没有证据块）—— 一并处理是
        # 「少一处就少一次误伤」的反面：动它没有任何收益，只会多一条可能改错的路径。
        raw_names = [n for n in zf.namelist() if n.endswith('-转录原文.md')]
        assert raw_names, zf.namelist()
        raw_text = zf.read(raw_names[0]).decode('utf-8')
    assert '[证据 1]' in text
    assert '## 证据附录' in text
    assert 'limo-block' not in text
    assert '证据附录' not in raw_text


def test_无证据块的文档导出不受线性化影响():
    """零回归：`tr_blk001`（只有 key-takeaway）的导出与改造前一致 —— 强调块照常
    包成 `md-block`、正文里没有多出一节「证据附录」。"""
    out = _only(_export(formats=['md']))
    assert '## 证据附录' not in out
    assert '核心结论' in out
    html = _only(_export(formats=['html']))
    assert 'class="md-block" data-block-type="key-takeaway"' in _body(html)
    assert '证据附录' not in _body(html)


def teardown_module(module):
    for d in _TMP:
        shutil.rmtree(d, ignore_errors=True)
