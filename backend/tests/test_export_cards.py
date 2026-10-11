"""卡片骨架的**导出跟随** —— HTML / Word 两条链。

本文件要防的失败全是**静默的**（本项目点名的「断言全过、肉眼等于没有」那一类）：

  A. `<html>` 上只写了 `data-card-style` 而漏了 `data-skeleton="card"`
     → `cards.css` 里所有 `[data-skeleton="card"] …` 规则**永不命中**
     → 导出的是**没有样式的卡片**，而「section 数量与 H2 一致」「CSS 里有卡片规则」
     这类断言**照样全绿**。
  B. Word 侧漏加 `section` 分支 → 落进 `else`「未知块：递归子元素」→ **内容一个不少**
     → 断「内容还在」照样绿。故必须断到「**被单单元格表格包裹 / 底纹生效**」这一层。
  C. `paper` 模式悄悄带上卡片属性 → 既有用户的导出观感被静默改动、零报错零提示。
  D. `_load_card_css` 读不到时静默返回空串 → 读者拿到一张张无边框无底纹的白块。
"""
from __future__ import annotations

import io
import sys

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.core.errors import AppError
from app.routers import export
from app.routers.export import _build_selfcontained_html, _load_card_css, _parse_card_vars
from app.services.docx_export import FILL_CARD, markdown_to_docx

#: ⚠️ 2026-10-11 起骨架轴由**标记**驱动：只有被 `card` 节修饰符标过的节才成卡。
#: 故「卡片」类用例要用 `CARD_MD`；`MD` 保留为**无标记**版本，供「零回归 / 默认
#: paper 逐字节不变」类用例使用（那些用例的输入必须真的不带标记，否则测的不是同一件事）。
MD = "开场白。\n\n## 甲\n内容一\n\n---\n\n## 乙\n内容二\n\n## 丙\n内容三\n"

#: 甲 / 乙标了 `card`、**丙没标** —— 混排是现在的常态，而「标了的成卡、没标的保持
#: 普通正文」这条判据必须由同一份样本同时覆盖两侧（只测「全标」的话，把实现改成
#: 「凡是 card 模式一律成卡」它照样绿）。
CARD_MD = (
    "开场白。\n\n"
    '## 甲\n<!-- limo-block:v1 type="card" -->\n内容一\n\n---\n\n'
    '## 乙\n<!-- limo-block:v1 type="card" -->\n内容二\n\n'
    "## 丙\n内容三\n"
)


# ═══════════════════════════════════════════════════════════════
# A. HTML 链
# ═══════════════════════════════════════════════════════════════

def test_默认_paper_导出与改造前一致():
    """零回归：不传参 = 今天的产物。三样东西一个都不能出现。"""
    html = _build_selfcontained_html(MD, "标题")
    assert "data-skeleton" not in html
    assert "data-card-style" not in html
    assert "md-card" not in html
    assert "md-cards" not in html
    # 内容一字不少
    for chunk in ("开场白。", "甲", "内容一", "乙", "内容二", "丙", "内容三"):
        assert chunk in html


def test_card_模式下_html_标签同时带两个属性():
    """对照 A：漏 `data-skeleton="card"` 时本条红 —— 而它是卡片规则的主作用域，
    只靠 `data-card-style` 的话共享的 CSS 在导出页上一条都不命中。"""
    html = _build_selfcontained_html(MD, "标题", skeleton="card", card_style="ink")
    assert 'data-skeleton="card"' in html
    assert 'data-card-style="ink"' in html


def test_card_模式下卡片数等于被标记的节数():
    """⚠️ 数量断的是**标记数（2）**而不是 H2 数（3）—— 后者是本轮改掉的那条无差别规则。
    同时必须断「未标记的丙节内容仍在」：把 `card` 判据写反（未标记的成卡、标记的不成卡）
    时，数量仍是 2，只有这条内容侧断言 + 下面那条「丙不在 section 里」能分辨。"""
    html = _build_selfcontained_html(CARD_MD, "标题", skeleton="card")
    assert html.count('<section class="md-card">') == 2
    for chunk in ("开场白。", "内容一", "内容二", "内容三"):
        assert chunk in html
    # 未标记的丙节**不在**任何 section 里（它原样穿过渲染链）
    tail = html[html.rindex("</section>") + len("</section>"):]
    assert "内容三" in tail and "<section" not in tail


def test_card_模式下标记行不落进产物():
    """节修饰符是内部协议 —— 读者不该在页面上看到 `<!-- limo-block … -->`。
    剥它的地方是渲染层（`parse_blocks` 的标记行剥除），切分器只管认不认得出它。

    ⚠️ 判据是**注释形式的标记行**，不是裸子串 `limo-block` —— 内联的 `cards.css`
    注释里就写着这个词（那是给读源码的人看的，不会显示给读者），拿裸子串去断会恒红。
    """
    html = _build_selfcontained_html(CARD_MD, "标题", skeleton="card")
    # ⚠️ 判据限定在 `<body>` 之后的**正文区**：内联的 `cards.css` 注释里就写着这句
    # 标记的示例（那是给读源码的人看的，不会显示给读者），拿整份 HTML 去断会恒红。
    body = html[html.index("<body"):]
    assert "limo-block" not in body
    assert 'type="card"' not in body


def test_card_模式下卡片_CSS_真的被内联():
    """⚠️ 断的是「导出页里出现了 `cards.css` 的选择器」，不是「文件存在」——
    文件存在但没拼进 style 块，效果与没有一样。"""
    # 注入 CSS 的门是 `skeleton == "card"`，与「有没有标记」无关 —— 用无标记的 MD
    # 正是不让这一条被标记干扰（有标记的那份由上面的用例管）。
    html = _build_selfcontained_html(MD, "标题", skeleton="card")
    css = _load_card_css()
    assert '[data-skeleton="card"]' in html
    assert '[data-skeleton="card"]' in css      # 自证：这份 CSS 确实是卡片规则那份
    assert "--card-radius" in html


def test_卡片_CSS_引用的每个界面令牌在导出页里都有定义():
    """**成对常量的守卫**：`cards.css` 里每处 `var(--tk-…)` 都必须在导出 HTML 里有定义。

    ⚠️ 导出主题（`HTML_THEMES`）是自包含的独立 CSS，**一个 `--tk-*` 都不提供** —— 而
    `var()` 无 fallback 时整条声明作废。漏一个的表现是「卡片少一块样式」，不报错、
    不抛异常，且**应用内预览里是好的**（那边 `--tk-*` 由主题系统提供）。这条是唯一
    能自动发现「cards.css 新引了一个令牌而导出侧没补」的地方。
    """
    import re as _re
    css = _re.sub(r"/\*[\s\S]*?\*/", "", _load_card_css())   # 剥注释：注释里提到名字不算引用
    used = sorted(set(_re.findall(r"var\((--tk-[a-z-]+)", css)))
    assert used, "cards.css 没引用任何 --tk-* —— 判据本身失效了（是不是改成别的命名了？）"
    html = _build_selfcontained_html(MD, "标题", skeleton="card")
    missing = [t for t in used if f"{t}:" not in html]
    assert missing == [], f"导出页缺少这些令牌的定义：{missing}"


def test_令牌三态齐全_暗色与打印段各有一套():
    """`@media (prefers-color-scheme: dark)` 与 `@media print` **会同时匹配**（既有教训），
    故导出页的令牌必须三态各自给全 —— 暗色段漏一个，那条声明在暗色下静默失效。"""
    html = _build_selfcontained_html(MD, "标题", skeleton="card")
    start = html.index("--tk-surface:")
    dark = html.index("@media (prefers-color-scheme: dark)", start)
    printing = html.index("@media print", dark)
    dark_block = html[dark:printing]
    print_block = html[printing:html.index("</style>", printing)]
    for tok in ("--tk-surface", "--tk-elevated", "--tk-hover", "--tk-border",
                "--tk-border-light", "--tk-foreground", "--tk-dimmed", "--tk-accent"):
        assert f"{tok}:" in dark_block, f"暗色段缺 {tok}"
        assert f"{tok}:" in print_block, f"打印段缺 {tok}"


def test_无_H2_时不产出卡片容器且内容不丢():
    """⚠️ 断的是**开标签字面量**而非 `"md-cards"` 子串 —— 后者在内联的 CSS 里
    （`.md-cards { display: flex }`）必然出现，那样断言恒假。"""
    html = _build_selfcontained_html("只有一段话\n第二段\n", "标题", skeleton="card")
    assert '<div class="md-cards">' not in html
    assert '<section class="md-card">' not in html
    assert "只有一段话" in html and "第二段" in html


def test_章末的_hr_成卡的剥掉_不成卡的保留():
    """三个方向各一条，缺一不可（2026-10-11 起判据挂在 `card` 上）：

    · **成卡的节**：章节边界已由卡的下边界承担，那条 `---` 挂在卡底只是噪音 → 剥；
    · **未标记的节**：它没有卡边界，`---` 是**唯一**的分隔手段 → 留。
      （漏了这条，「开了卡片模式、普通章节之间的分隔线却消失了」—— 与骨架轴无关的损失。）
    · **章节内部**的 `---`：在分隔话题，不是章节边界 → 恒留。
    """
    html = _build_selfcontained_html(CARD_MD, "标题", skeleton="card")
    assert html.count("<hr") == 0   # 甲节末尾那条随成卡被剥
    plain = _build_selfcontained_html(
        "## 甲\n内容一\n\n---\n\n## 乙\n内容二\n", "t", skeleton="card")
    assert plain.count("<hr") == 1  # 都未标记 → 保留
    inner = _build_selfcontained_html(
        '## 甲\n<!-- limo-block:v1 type="card" -->\n话题一\n\n---\n\n话题二\n',
        "t", skeleton="card")
    assert inner.count("<hr") == 1  # 成卡也留章内的那条


def test_非法视觉方案静默回退_gpt():
    """视觉方案错了不改变产物的**性质**（判据同合成管道的维度清洗），不该让整次导出失败。"""
    html = _build_selfcontained_html(MD, "标题", skeleton="card", card_style="不存在的皮肤")
    assert 'data-card-style="gpt"' in html


def test_load_card_css_读到的是真文件():
    css = _load_card_css()
    assert ".md-card" in css
    assert "--card-bg-base" in css


def test_显式传_paper_与不传参数逐字节等价():
    """「默认值就是 paper」这条契约本身要能被证伪 —— 不传即走用户今天拿到的那条路。"""
    assert (_build_selfcontained_html(MD, "标题")
            == _build_selfcontained_html(MD, "标题", skeleton="paper", card_style="ink"))


def test_打包路径_MEIPASS_下也读得到(tmp_path, monkeypatch):
    """⚠️ 这条覆盖的是**开发环境永远走不到**的分支：`sys._MEIPASS` 只在 PyInstaller
    打包后存在。不测它等于那几行从没被执行过 —— 真出问题是在用户机器上（500），
    而开发机永远绿。"""
    (tmp_path / "styles").mkdir()
    (tmp_path / "styles" / "cards.css").write_text("/* packed */ .md-card{}", encoding="utf-8")
    # 让开发路径那条候选**落空**：它指向仓库里的真文件，不挪开就永远先命中，
    # 第二个候选也就永远测不到。`__file__` 是模块全局，patch 它即可。
    monkeypatch.setattr(export, "__file__", str(tmp_path / "app" / "routers" / "export.py"))
    monkeypatch.setattr(sys, "_MEIPASS", str(tmp_path), raising=False)
    assert _load_card_css() == "/* packed */ .md-card{}"


def test_读不到_cards_css_时显式抛错而非静默降级(monkeypatch):
    """对照 D：返回空串的话导出照样「成功」，只是卡片一点样式都没有。"""

    class _AlwaysMissing:
        """假的 `Path`：支持 `Path(...).resolve().parents[n] / "a" / "b".is_file()`。"""

        def __init__(self, *a, **kw):
            pass

        def resolve(self):
            return self

        @property
        def parents(self):
            return [self] * 8

        def __truediv__(self, _other):
            return self

        def is_file(self):
            return False

        def __str__(self):
            return "<fake-path>"

    monkeypatch.setattr(export, "Path", _AlwaysMissing)
    with pytest.raises(AppError) as ei:
        _load_card_css()
    assert ei.value.code == "EXPORT_CARD_CSS_MISSING"
    # 部署期错误，重试不会好转 —— 不可恢复
    assert ei.value.recoverable is False


# ═══════════════════════════════════════════════════════════════
# C. 逐维度微调（`card_vars`，2026-10-11）
# ═══════════════════════════════════════════════════════════════

#: 前端 `serializeCardVars` 会发出来的真实形态（`lib/card-prefs.ts`）。
_LEGAL_VARS = ("--card-bg-base:var(--tk-hover);--card-lift:40%"
               ";--card-border:2px solid var(--tk-border)")


def test_微调_合法串解析成_root_段():
    css = _parse_card_vars(_LEGAL_VARS)
    assert css.startswith(":root{") and css.endswith("}")
    assert "--card-bg-base:var(--tk-hover)" in css
    assert "--card-lift:40%" in css
    assert "--card-border:2px solid var(--tk-border)" in css


def test_微调_空串给空串():
    """没调过的用户走今天这条逐字未改的分支 —— `_build_selfcontained_html` 据此
    **一个字节都不多**（下面还有一条逐字节等价的用例钉住）。"""
    assert _parse_card_vars("") == ""
    assert _parse_card_vars(";;;") == ""


def test_微调_不认识的名字被丢弃():
    """白名单外的变量一律不写 —— 否则 query 参数就成了「往导出页注入任意 CSS 变量」
    的通道（`--tk-border:#000` 之类能整页改色）。"""
    assert _parse_card_vars("--tk-border:red") == ""
    assert _parse_card_vars("--card-radius:20px") == ""   # 圆角刻意没开成旋钮
    # 逐项丢弃：一项非法**不影响**同串里的合法项
    css = _parse_card_vars("--tk-border:red;--card-lift:40%")
    assert "--tk-border" not in css and "--card-lift:40%" in css


@pytest.mark.parametrize("evil", [
    "--card-lift:0}",                        # 提前闭合规则块
    "--card-lift:40%;--card-radius:20px",    # 用分号续写
    "--card-lift:</style><script>alert(1)</script>",
    "--card-border:1px solid red/*",         # 注释起始
    "--card-bg-base:url(//evil.example/x)",  # 外链（`/` 与 `*` 都不在白名单里）
])
def test_微调_值里的逃逸字符全被拦下(evil):
    """对照：把白名单放宽成「只要不含分号」→ 这几条至少有一条能逃出去。
    ⚠️ 判据是**整串没有可疑字符**，不是「解析结果非空」—— 后者对逃逸串也成立。"""
    css = _parse_card_vars(evil)
    # 整串只该有 `:root{` 这一对花括号 —— 多一个就是「值里逃出了一个规则块」
    assert css.count("{") <= 1 and css.count("}") <= 1, css
    for bad in ("<", ">", "/", "*", '"', "'", "\\"):
        assert bad not in css, f"{bad!r} 混进去了：{css}"


def test_微调_必须排在_cards_css_之后():
    """⚠️ 顺序即优先级：方案块 `[data-card-style="x"]` 与微调的 `:root` **特异性相同**
    （都是 0,1,0），同特异性下后者胜。放到前面会被方案原样压住 ——
    症状是「预览里调过的项，导出还是方案取值」，而导出成功、不报错。
    ⚠️ 用 `--sh-0:` 定位 cards.css（投影等级层，只在那一份里出现）。"""
    html = _build_selfcontained_html(MD, "标题", skeleton="card", card_vars="--card-lift:40%")
    assert "--card-lift:40%" in html
    assert html.index("--card-lift:40%") > html.index("--sh-0:"), "微调段跑到 cards.css 前面了"


def test_微调_没传时与改造前逐字节等价():
    """零回归：`card_vars` 默认空串 ⇒ 多出来的那个 `<style>` 根本不存在。"""
    assert (_build_selfcontained_html(MD, "标题", skeleton="card")
            == _build_selfcontained_html(MD, "标题", skeleton="card", card_vars=""))


def test_微调_paper_模式下不注入():
    """`card_vars` 只在骨架为 `card` 时消费 —— `paper` 下没有卡片，
    注入了也是死规则，还会破坏「paper 导出逐字节不变」。"""
    html = _build_selfcontained_html(MD, "标题", card_vars="--card-lift:40%")
    assert "--card-lift:40%" not in html


# ═══════════════════════════════════════════════════════════════
# B. Word 链
# ═══════════════════════════════════════════════════════════════

def _tables(data: bytes) -> list:
    return Document(io.BytesIO(data)).tables


def test_word_paper_默认不产生卡片表格():
    assert len(_tables(markdown_to_docx(MD, "标题"))) == 0


def test_word_card_每节包成单单元格表格():
    """对照 B：断在**结构层**（表格数 / 单元格数），不断「内容还在」——
    漏加 section 分支时内容由 `else` 兜底照样渲染，那种断言绿得毫无意义。

    ⚠️ 数量是**标记数（2）**：丙节没标 `card`，它不该被包成表格。"""
    data = markdown_to_docx(CARD_MD, "标题", skeleton="card")
    tables = _tables(data)
    assert len(tables) == 2
    for t in tables:
        assert len(t.rows) == 1
        assert len(t.columns) == 1


def test_word_卡片底纹真的应用了():
    """结构层再深一层：`w:tcPr/w:shd` 存在且取值正确（只加表格不加底纹 = 白框）。"""
    tables = _tables(markdown_to_docx(CARD_MD, "标题", skeleton="card"))
    tc = tables[0].rows[0].cells[0]._tc
    tcPr = tc.find(qn("w:tcPr"))
    assert tcPr is not None
    shd = tcPr.find(qn("w:shd"))
    assert shd is not None
    assert shd.get(qn("w:fill")) == FILL_CARD


def test_word_卡片内标题是_Heading_样式而不是普通段落():
    """`_CellProxy` 存在的唯一理由就是 `_Cell` 没有 `add_heading` ——
    这条钉住它真的被用上了（掉了的话标题会退化成正文段落，观感全变、零报错）。"""
    tables = _tables(markdown_to_docx(
        '## 甲\n<!-- limo-block:v1 type="card" -->\n正文\n', "标题", skeleton="card"))
    styles = {p.style.name for p in tables[0].rows[0].cells[0].paragraphs}
    assert any(s.startswith("Heading") for s in styles), styles


def test_word_对照_裸_section_标签不会变卡片():
    """自证有区分力：表格来自 `class="md-card"` 这个**判据**，不是来自 `section` 标签本身。
    把 md 里手写一个不带 class 的 section，走的是 `else` 递归 → 0 表格。"""
    data = markdown_to_docx("<section>\n\n## 甲\n正文\n\n</section>\n", "t")
    assert len(_tables(data)) == 0


def test_word_卡片内容零增减():
    """结构对了还要内容对 —— 两张卡不能把正文吞掉或复制一份，**未标记的丙节也不能丢**
    （它的内容不在表格里，故要另取一次文档正文来断）。"""
    data = markdown_to_docx(CARD_MD, "标题", skeleton="card")
    got = "\n".join(
        p.text for t in _tables(data) for p in t.rows[0].cells[0].paragraphs
    )
    for chunk in ("甲", "内容一", "乙", "内容二"):
        assert got.count(chunk) == 1, f"{chunk} 出现 {got.count(chunk)} 次"
    body = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert body.count("丙") == 1 and body.count("内容三") == 1
